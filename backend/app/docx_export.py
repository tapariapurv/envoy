"""Markdown position paper -> editable Word document (python-docx).

"envoy" mirrors the in-app paper preview: letterhead, numbered section badges, shaded pull-quote,
numbered proposal cards and highlighted key terms. "conference" is the plain Times New Roman 12 format
many conferences require. Both use real Word styles (Title, Heading 1, …) so the file stays editable.
"""
import io
import re
from html.parser import HTMLParser

import markdown
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SERIF, SANS = "Georgia", "Helvetica Neue"  # same fonts as the preview (.paper in globals.css)
C = {"ink": "1D1C1A", "muted": "75716A", "accent": "4F46E5", "tint": "F1F0FD", "badge": "E4E3FB", "mark": "E0DEFB",
     "card": "EFEDE8", "line": "E4E1DA", "warn": "C2410C", "warn_bg": "FBE9DC"}
STYLES = {
    "envoy": {"body": SERIF, "head": SERIF, "label": SANS, "size": 11, "line": Pt(18), "justify": False, "rich": True},
    "conference": {"body": "Times New Roman", "head": "Times New Roman", "label": "Times New Roman", "size": 12,
                   "line": 1.15, "justify": True, "rich": False},
}
STYLES["modern"], STYLES["classic"] = STYLES["envoy"], STYLES["conference"]  # older names
META_KEYS = ("committee", "topic", "country", "delegation", "delegate", "school", "conference")
META_RE = re.compile(r"^\s*(?:[-*]\s*)?\**\s*(%s)s?\s*(?::\**|\**:)\s*(.*?)\s*$" % "|".join(META_KEYS), re.I)


def split_meta(md: str) -> tuple[str, dict, str]:
    """Pull a leading '# Title' and '**Committee:** X' style lines off the top of the paper."""
    title, meta, lines = "", {}, md.strip().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].strip()
        if not ln or ln in ("---", "***"):
            i += 1
        elif ln.startswith("# ") and not title and not meta:
            title = ln[2:].strip().strip("*")
            i += 1
        elif m := META_RE.match(ln):
            if v := m[2].strip().strip("*").strip():  # "**Delegate:**" with no value is dropped
                meta[m[1].lower().replace("delegation", "country")] = v
            i += 1
        else:
            break
    return title, meta, "\n".join(lines[i:])


# ---------- OOXML helpers ----------
def _rgb(c: str) -> RGBColor:
    return RGBColor.from_string(c)


def _el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def _shade(target, fill: str):
    """Background fill for a run, paragraph or table cell."""
    props = target._r.get_or_add_rPr() if hasattr(target, "_r") else (
        target._tc.get_or_add_tcPr() if hasattr(target, "_tc") else target._p.get_or_add_pPr())
    props.append(_el("w:shd", val="clear", color="auto", fill=fill))


def _border(paragraph, side: str, color: str, size: int = 12, space: int = 4):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        ppr.append(bdr := OxmlElement("w:pBdr"))
    bdr.append(_el(f"w:{side}", val="single", sz=size, space=space, color=color))


def _cell_style(cell, fill: str | None, border: str | None, margins=(140, 200, 140, 200)):
    tcpr = cell._tc.get_or_add_tcPr()
    if fill:
        _shade(cell, fill)
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b.append(_el(f"w:{side}", val="single" if border else "nil", sz=4, color=border or "auto"))
    tcpr.append(b)
    m = OxmlElement("w:tcMar")
    for side, v in zip(("top", "left", "bottom", "right"), margins):
        m.append(_el(f"w:{side}", w=v, type="dxa"))
    tcpr.append(m)


def _table(doc, cols: list[float]):
    t = doc.add_table(rows=1, cols=len(cols))
    t.autofit = False
    t._tbl.tblPr.append(_el("w:tblBorders"))  # no default grid
    t._tbl.tblPr.append(_el("w:tblLayout", type="fixed"))
    for col, cell, w in zip(t.columns, t.rows[0].cells, cols):
        col.width = cell.width = Cm(w)
    return t


def _field(run, instr: str):
    """Live Word field (e.g. PAGE) so numbers update in Word."""
    run._r.append(_el("w:fldChar", fldCharType="begin"))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    run._r.append(it)
    run._r.append(_el("w:fldChar", fldCharType="separate"))
    run._r.append(_el("w:fldChar", fldCharType="end"))


def _font(run, name: str, size: float | None = None, color: str | None = None, bold=None, italic=None):
    run.font.name = name
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = _rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def _setup(doc: Document, st: dict):
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        sec.left_margin = sec.right_margin = Cm(2.4)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = st["body"], Pt(st["size"]), _rgb(C["ink"])
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), st["body"])
    pf = normal.paragraph_format
    pf.space_after, pf.line_spacing = Pt(9), st["line"]
    if st["justify"]:
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def style(name, font, size, color, bold=False, italic=False, before=0, after=6):
        s = doc.styles[name]
        s.font.name, s.font.size, s.font.bold, s.font.italic, s.font.color.rgb = font, Pt(size), bold, italic, _rgb(color)
        rpr = s.element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set(qn("w:ascii"), font)
        rpr.get_or_add_rFonts().set(qn("w:hAnsi"), font)
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
        s.paragraph_format.line_spacing = 1.1
        s.paragraph_format.keep_with_next = True
        return s

    rich, base = st["rich"], st["size"]
    style("Title", st["head"], 30 if rich else base + 12, C["ink"], bold=not rich, after=8)
    style("Heading 1", st["head"], 17 if rich else base + 2, C["ink"], bold=not rich, before=18, after=8)
    style("Heading 2", st["head"], 13 if rich else base, C["ink"], bold=True, before=12, after=4)
    style("Heading 3", st["head"], base, C["muted"], bold=True, italic=True, before=8, after=3)
    tp = doc.styles["Title"].element.get_or_add_pPr()
    for b in tp.findall(qn("w:pBdr")):  # default Title has its own rule; we draw ours
        tp.remove(b)


class _Builder(HTMLParser):
    """Walks Markdown-rendered HTML and emits python-docx paragraphs, runs and tables."""

    def __init__(self, doc, st):
        super().__init__(convert_charrefs=True)
        self.doc, self.st, self.rich = doc, st, st["rich"]
        self.container, self.fresh = doc, False  # where paragraphs go (document or a table cell)
        self.p = None
        self.fmt = {"b": 0, "i": 0, "code": 0, "a": 0}
        self.lists: list[str] = []
        self.card_n = 0
        self.in_card = self.in_quote = False
        self.sec = 0
        self.table: list[list[str]] | None = None
        self.cell: list[str] | None = None

    def _new(self, style=None):
        if self.fresh:  # a new table cell already holds one empty paragraph
            self.p, self.fresh = self.container.paragraphs[0], False
            if style:
                self.p.style = self.doc.styles[style]
        else:
            self.p = self.container.add_paragraph(style=style)
        return self.p

    def _enter_cell(self, cell):
        self.container, self.fresh = cell, True

    def _exit_cell(self):
        self.container, self.fresh, self.p = self.doc, False, None
        gap = self.doc.add_paragraph()
        gap.paragraph_format.space_after, gap.paragraph_format.line_spacing = Pt(0), Pt(7)

    # --- tags ---
    def handle_starttag(self, tag, attrs):
        if self.table is not None:
            if tag == "tr":
                self.table.append([])
            elif tag in ("td", "th"):
                self.cell = []
            return
        if tag in ("h1", "h2"):
            self.sec += 1
            p = self._new("Heading 1")
            if self.rich:
                _shade(_font(p.add_run(f" {self.sec:02d} "), SANS, 8, C["accent"], bold=True), C["badge"])
                p.add_run("  ")
        elif tag == "h3":
            self._new("Heading 2")
        elif tag in ("h4", "h5", "h6"):
            self._new("Heading 3")
        elif tag == "p":
            if not (self.p is not None and self.p.text.strip("“ ") == "" and (self.in_card or self.in_quote or self.lists)):
                self._new()
        elif tag in ("ul", "ol"):
            self.lists.append(tag)
        elif tag == "li":
            if self.rich and self.lists == ["ol"]:  # top-level numbered list -> proposal cards
                self.card_n += 1
                t = _table(self.doc, [0.9, 15.1])
                num, body = t.rows[0].cells
                _cell_style(num, None, None, (150, 0, 0, 0))
                num.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                np = num.paragraphs[0]
                np.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _shade(_font(np.add_run(f" {self.card_n} "), SANS, 9, "FFFFFF", bold=True), C["accent"])
                _cell_style(body, C["card"], C["line"])
                self._enter_cell(body)
                self.in_card = True
                self._new().paragraph_format.space_after = Pt(0)
            else:
                kind = "List Number" if self.lists and self.lists[-1] == "ol" else "List Bullet"
                self._new(kind if len(self.lists) == 1 else f"{kind} 2").paragraph_format.space_after = Pt(4)
        elif tag == "blockquote":
            if self.rich:
                t = _table(self.doc, [16.0])
                _cell_style(t.rows[0].cells[0], C["tint"], None, (200, 320, 200, 320))
                self._enter_cell(t.rows[0].cells[0])
                self.in_quote = True
                q = self._new()
                q.paragraph_format.space_after, q.paragraph_format.line_spacing = Pt(0), 1.25
                _font(q.add_run("“ "), SERIF, 26, C["accent"])
            else:
                self._new().paragraph_format.left_indent = Cm(1)
                self.fmt["i"] += 1
        elif tag in ("strong", "b"):
            self.fmt["b"] += 1
        elif tag in ("em", "i"):
            self.fmt["i"] += 1
        elif tag == "code":
            self.fmt["code"] += 1
        elif tag == "a":
            self.fmt["a"] += 1
        elif tag == "br" and self.p is not None:
            self.p.add_run().add_break(WD_BREAK.LINE)
        elif tag == "hr":
            _border(self._new(), "bottom", C["line"], size=6)
        elif tag == "table":
            self.table = []

    def handle_endtag(self, tag):
        if self.table is not None:
            if tag in ("td", "th") and self.cell is not None:
                self.table[-1].append("".join(self.cell).strip())
                self.cell = None
            elif tag == "table":
                self._emit_table(self.table)
                self.table = None
            return
        if tag in ("ul", "ol") and self.lists:
            self.lists.pop()
        elif tag == "li" and self.in_card and self.lists == ["ol"]:
            self.in_card = False
            self._exit_cell()
        elif tag == "blockquote":
            if self.in_quote:
                self.in_quote = False
                self._exit_cell()
            else:
                self.fmt["i"] -= 1
        elif tag in ("strong", "b"):
            self.fmt["b"] -= 1
        elif tag in ("em", "i"):
            self.fmt["i"] -= 1
        elif tag == "code":
            self.fmt["code"] -= 1
        elif tag == "a":
            self.fmt["a"] -= 1
        elif tag in ("p", "h1", "h2", "h3", "h4", "h5", "h6") and not (self.in_card or self.in_quote):
            self.p = None

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)
            return
        if not data.strip() and ("\n" in data or self.p is None):  # formatting whitespace between block tags
            return
        p = self.p if self.p is not None else self._new()
        data = re.sub(r"(?<!-)--(?!-)", "—", re.sub(r"(\S) - ", "\\1 – ", data))  # typographic dashes
        for part in re.split(r"(\[citation needed\]|\[\d{1,2}(?:,\s*\d{1,2})*\])", data, flags=re.I):
            if not part:
                continue
            if part.lower() == "[citation needed]":
                r = _font(p.add_run(" citation needed "), SANS if self.rich else self.st["body"], 8, C["warn"], italic=True)
                if self.rich:
                    _shade(r, C["warn_bg"])
                continue
            if re.fullmatch(r"\[[\d,\s]+\]", part):
                r = p.add_run(part[1:-1] if self.rich else part)
                r.font.superscript = True
                r.font.color.rgb = _rgb(C["accent"] if self.rich else C["muted"])
                r.bold = self.rich or None
                continue
            r = p.add_run(part)
            if self.in_quote:
                _font(r, SERIF, 14, C["ink"], italic=True)
            r.bold = bool(self.fmt["b"]) or None
            r.italic = bool(self.fmt["i"]) or self.in_quote or None
            if self.fmt["b"] and self.rich:
                if self.in_card:
                    r.font.color.rgb = _rgb(C["accent"])
                elif not self.in_quote and p.style.name == "Normal":
                    _shade(r, C["mark"])
            if self.fmt["code"]:
                _font(r, "Menlo", self.st["size"] - 1.5)
            if self.fmt["a"]:
                r.font.color.rgb, r.underline = _rgb(C["accent"]), True

    def _emit_table(self, rows):
        rows = [r for r in rows if r]
        if not rows:
            return
        t = self.doc.add_table(rows=len(rows), cols=max(map(len, rows)))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                cell = t.cell(i, j)
                cell.text = v
                if i == 0:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
        self.doc.add_paragraph()


def build(markdown_text: str, profile: dict, style: str = "envoy", fallback_title: str = "Position Paper") -> bytes:
    st = STYLES.get(style, STYLES["envoy"])
    rich = st["rich"]
    title, meta, body = split_meta(markdown_text)
    meta = {"country": profile.get("delegate_country", ""), "committee": profile.get("committee", ""),
            "topic": profile.get("topic", ""), **{k: v for k, v in meta.items() if v}}

    doc = Document()
    _setup(doc, st)
    doc.core_properties.title = title or fallback_title
    doc.core_properties.subject = meta.get("topic", "")
    doc.core_properties.author = meta.get("delegate") or meta.get("country") or "Envoy"
    doc.core_properties.comments = "Created with Envoy"

    # --- letterhead (mirrors .paper-head) ---
    label = doc.add_paragraph()
    label.paragraph_format.space_after = Pt(4)
    _font(label.add_run((title if title and title.lower() != "position paper" else "Position Paper").upper()),
          st["label"], 8 if rich else 10, C["accent"] if rich else C["ink"], bold=True)
    doc.add_paragraph(meta.get("country") or title or fallback_title, style="Title")
    fields = [(k, meta.get(k.lower())) for k in ("Committee", "Topic", "Delegate", "School", "Conference")]
    info = doc.add_paragraph()
    info.paragraph_format.space_after, info.paragraph_format.line_spacing = Pt(16), 1.6
    for k, v in (f for f in fields if f[1]):
        if rich:  # pill-like chips
            _shade(_font(info.add_run(f"  {v}  "), SANS, 8.5, C["ink"]), C["card"])
            info.add_run(" ")
        else:
            if info.runs:
                info.add_run().add_break(WD_BREAK.LINE)
            _font(info.add_run(f"{k}: "), st["body"], st["size"], bold=True)
            info.add_run(v)
    _border(info, "bottom", C["accent"] if rich else C["ink"], size=18 if rich else 8, space=10)

    # --- body ---
    b = _Builder(doc, st)
    b.feed(markdown.markdown(body, extensions=["tables", "sane_lists"]))
    b.close()

    # --- footer page number (live field) ---
    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(_font(fp.add_run(), SANS if rich else st["body"], 8.5, C["muted"]), "PAGE")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
